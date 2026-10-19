from __future__ import annotations

import base64
import html
import math
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / "member3_assets"
OUTPUT = ROOT / "Thanh_vien_3.docx"
HTML_OUTPUT = ROOT / "member3_content.html"
ASSETS.mkdir(exist_ok=True)

FONT_REGULAR = Path(r"C:\Windows\Fonts\arial.ttf")
FONT_BOLD = Path(r"C:\Windows\Fonts\arialbd.ttf")

NAVY = "123B5D"
TEAL = "0E7490"
PALE_BLUE = "EAF4F8"
PALE_GRAY = "F5F7F9"
GRAY = "D9E0E5"
GREEN = "2E7D32"
RED = "B3261E"
ORANGE = "D97706"
BLACK = "111827"
WHITE = "FFFFFF"


def font(size: int, bold: bool = False):
    return ImageFont.truetype(str(FONT_BOLD if bold else FONT_REGULAR), size)


def wrap(draw: ImageDraw.ImageDraw, text: str, fnt, max_width: int) -> list[str]:
    words = text.split()
    lines: list[str] = []
    current = ""
    for word in words:
        test = word if not current else current + " " + word
        if draw.textbbox((0, 0), test, font=fnt)[2] <= max_width:
            current = test
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines


def text_center(draw, box, text, fnt, fill=BLACK, line_gap=8):
    x1, y1, x2, y2 = box
    lines = wrap(draw, text, fnt, x2 - x1 - 28)
    heights = [draw.textbbox((0, 0), line, font=fnt)[3] for line in lines]
    total = sum(heights) + line_gap * max(0, len(lines) - 1)
    y = y1 + (y2 - y1 - total) / 2
    for line, h in zip(lines, heights):
        w = draw.textbbox((0, 0), line, font=fnt)[2]
        draw.text((x1 + (x2 - x1 - w) / 2, y), line, font=fnt, fill="#" + fill)
        y += h + line_gap


def box(draw, rect, text, fill=PALE_BLUE, outline=NAVY, bold=False, size=30, radius=16):
    draw.rounded_rectangle(rect, radius=radius, fill="#" + fill, outline="#" + outline, width=4)
    text_center(draw, rect, text, font(size, bold), BLACK)


def arrow(draw, start, end, color=NAVY, width=5, label=None, label_offset=(0, -28)):
    draw.line([start, end], fill="#" + color, width=width)
    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    length = 18
    a1 = angle + math.pi * 0.82
    a2 = angle - math.pi * 0.82
    p1 = (end[0] + length * math.cos(a1), end[1] + length * math.sin(a1))
    p2 = (end[0] + length * math.cos(a2), end[1] + length * math.sin(a2))
    draw.polygon([end, p1, p2], fill="#" + color)
    if label:
        fnt = font(23)
        w = draw.textbbox((0, 0), label, font=fnt)[2]
        mx = (start[0] + end[0]) / 2 + label_offset[0]
        my = (start[1] + end[1]) / 2 + label_offset[1]
        draw.rectangle((mx - w / 2 - 8, my - 4, mx + w / 2 + 8, my + 28), fill="#FFFFFF")
        draw.text((mx - w / 2, my), label, font=fnt, fill="#" + BLACK)


def canvas(title: str, size=(1800, 1120)):
    img = Image.new("RGB", size, "white")
    draw = ImageDraw.Draw(img)
    draw.text((70, 45), title, font=font(42, True), fill="#" + NAVY)
    draw.line((70, 110, size[0] - 70, 110), fill="#" + TEAL, width=5)
    return img, draw


def save_flow():
    img, d = canvas("Luồng xử lý một trận đấu")
    coords = [
        (80, 170, 400, 290, "Chấp nhận thách đấu"),
        (520, 170, 840, 290, "Tạo GameRoom"),
        (960, 170, 1280, 290, "Gán vai ngẫu nhiên"),
        (1400, 170, 1720, 290, "Bắt đầu vòng"),
        (1400, 430, 1720, 550, "Nhận lựa chọn trong 10 giây"),
        (960, 430, 1280, 550, "So sánh và cộng điểm"),
        (520, 430, 840, 550, "Đổi vai và tăng số vòng"),
        (80, 430, 400, 550, "Đủ 6 vòng?"),
        (80, 700, 400, 820, "Hai người hòa?"),
        (520, 700, 840, 820, "Chơi cặp 2 vòng phụ"),
        (960, 700, 1280, 820, "Gửi kết quả cuối"),
        (1400, 700, 1720, 820, "Đóng phòng")
    ]
    for x1, y1, x2, y2, t in coords:
        fill = "FFF4E5" if "?" in t else PALE_BLUE
        box(d, (x1, y1, x2, y2), t, fill=fill, bold=True if t in ("Đủ 6 vòng?", "Hai người hòa?") else False)
    for a, b in [((400,230),(520,230)),((840,230),(960,230)),((1280,230),(1400,230)),((1560,290),(1560,430)),((1400,490),(1280,490)),((960,490),(840,490)),((520,490),(400,490)),((240,550),(240,700)),((400,760),(520,760)),((1280,760),(1400,760))]:
        arrow(d, a, b)
    arrow(d, (240, 550), (240, 700), label="Đủ", label_offset=(52, 0))
    arrow(d, (400, 760), (520, 760), label="Hòa", label_offset=(0, -36))
    d.line((240, 700, 240, 640, 1120, 640, 1120, 700), fill="#" + NAVY, width=5)
    arrow(d, (1120, 640), (1120, 700))
    d.text((870, 602), "Không hòa", font=font(23), fill="#" + BLACK)
    d.line((80, 490, 45, 490, 45, 340, 1560, 340, 1560, 430), fill="#" + NAVY, width=5)
    arrow(d, (1560, 340), (1560, 430))
    d.text((62, 365), "Chưa đủ 6 vòng", font=font(23), fill="#" + BLACK)
    d.line((680, 820, 680, 900, 240, 900, 240, 820), fill="#" + NAVY, width=5)
    arrow(d, (240, 900), (240, 820))
    d.text((355, 865), "Xét lại điểm sau cặp vòng", font=font(23), fill="#" + BLACK)
    d.text((70, 1010), "Ngoại lệ tại mọi thời điểm: timeout, Exit hoặc mất kết nối -> kết thúc theo luật xử lý sự cố.", font=font(25), fill="#" + RED)
    path = ASSETS / "01_luong_tran_dau.png"
    img.save(path, quality=95)
    return path


def save_classes():
    img, d = canvas("Cấu trúc các lớp phía Server")
    boxes = {
        "MatchManager": (670, 155, 1130, 275),
        "GameRoom": (670, 400, 1130, 540),
        "RoundManager": (120, 690, 500, 830),
        "GameEngine": (530, 690, 910, 830),
        "PuzzleManager": (940, 690, 1320, 830),
        "AirdropManager": (1350, 690, 1730, 830),
    }
    descriptions = {
        "MatchManager": "Tạo và kết thúc trận",
        "GameRoom": "Giữ trạng thái hai người chơi",
        "RoundManager": "Timer, vai trò, số vòng",
        "GameEngine": "Luật thắng thua và điểm",
        "PuzzleManager": "Câu đố và người trả lời sớm",
        "AirdropManager": "Vị trí ngẫu nhiên và buff",
    }
    for name, rect in boxes.items():
        d.rounded_rectangle(rect, radius=18, fill="#" + (PALE_BLUE if name != "GameRoom" else "DDF4EA"), outline="#" + (NAVY if name != "GameRoom" else GREEN), width=4)
        x1,y1,x2,y2 = rect
        d.text((x1+24,y1+22), name, font=font(31,True), fill="#"+NAVY)
        d.text((x1+24,y1+72), descriptions[name], font=font(24), fill="#"+BLACK)
    arrow(d,(900,275),(900,400),label="quản lý")
    starts=[(780,540),(850,540),(950,540),(1030,540)]
    ends=[(310,690),(720,690),(1130,690),(1540,690)]
    labels=["điều phối vòng","gọi luật","tạo câu đố","tạo Airdrop"]
    for s,e,l in zip(starts,ends,labels): arrow(d,s,e,label=l,label_offset=(0,-24))
    d.text((180, 930), "GameRoom là trung tâm trạng thái; các manager tách riêng từng trách nhiệm để dễ kiểm thử và tránh xử lý chồng chéo.", font=font(26), fill="#"+BLACK)
    path=ASSETS/"02_so_do_lop.png"; img.save(path, quality=95); return path


def save_states():
    img, d = canvas("Vòng đời của GameRoom")
    states = [
        ((100,190,430,310),"CREATED"),
        ((570,190,950,310),"PREPARING_ROUND"),
        ((1090,190,1470,310),"WAITING_MOVES"),
        ((1090,500,1470,620),"RESOLVING"),
        ((570,500,950,620),"ROUND_FINISHED"),
        ((100,500,430,620),"FINISHED"),
    ]
    for rect,name in states:
        fill="FCE8E6" if name=="FINISHED" else ("DDF4EA" if name=="ROUND_FINISHED" else PALE_BLUE)
        box(d,rect,name,fill=fill,bold=True,size=28)
    # Nhãn được đặt cạnh đường nối, không phủ lên đường để tránh cảm giác mũi tên bị cắt.
    arrow(d,(430,250),(570,250))
    d.text((445,198),"Đủ 2 người",font=font(22),fill="#"+BLACK)
    arrow(d,(950,250),(1090,250))
    d.text((965,198),"ROUND_START",font=font(22),fill="#"+BLACK)
    arrow(d,(1280,310),(1280,500))
    d.text((1310,390),"Đủ 2 lựa chọn",font=font(22),fill="#"+BLACK)
    d.text((1310,425),"hoặc hết 10 giây",font=font(22),fill="#"+BLACK)
    arrow(d,(1090,560),(950,560))
    d.text((955,505),"ROUND_RESULT",font=font(22),fill="#"+BLACK)
    arrow(d,(570,560),(430,560))
    d.text((360,505),"Đủ 6 vòng",font=font(22),fill="#"+BLACK)
    arrow(d,(760,500),(760,310))
    d.text((795,390),"Còn vòng",font=font(22),fill="#"+BLACK)
    box(d,(600,790,1600,930),"Exit hoặc mất kết nối: khóa phòng, xác định kết quả theo luật và chuyển sang FINISHED",fill="FFF4E5",outline=ORANGE,size=27)
    d.line((600,860,500,860,500,700,265,700),fill="#"+ORANGE,width=5)
    arrow(d,(265,700),(265,620),color=ORANGE)
    d.text((285,655),"Sự cố từ mọi trạng thái đang chạy",font=font(22),fill="#"+ORANGE)
    path=ASSETS/"03_trang_thai_gameroom.png"; img.save(path, quality=95); return path


def save_sequence():
    img, d = canvas("Trình tự xử lý lựa chọn trong một vòng")
    xs=[180,620,1060,1500]
    names=["Client A","ClientHandler","RoundManager","Client B"]
    for x,n in zip(xs,names):
        box(d,(x-140,150,x+140,240),n,fill=PALE_GRAY,bold=True,size=27)
        d.line((x,240,x,1010),fill="#8A99A6",width=3)
    def msg(y,a,b,label,color=NAVY,dashed=False):
        if dashed:
            step=20
            x1,x2=xs[a],xs[b]
            direction=1 if x2>x1 else -1
            x=x1
            while (x<x2 if direction>0 else x>x2):
                nx=x+direction*min(step,abs(x2-x))
                d.line((x,y,nx,y),fill="#"+color,width=4)
                x=nx+direction*10
            arrow(d,(x2-direction*24,y),(x2,y),color=color,width=4)
        else: arrow(d,(xs[a],y),(xs[b],y),color=color,width=4)
        f=font(23); w=d.textbbox((0,0),label,font=f)[2]
        d.rectangle(((xs[a]+xs[b])/2-w/2-8,y-36,(xs[a]+xs[b])/2+w/2+8,y-5),fill="white")
        d.text(((xs[a]+xs[b])/2-w/2,y-34),label,font=f,fill="#"+BLACK)
    msg(330,0,1,"MOVE;LEFT")
    msg(430,1,2,"validate + lưu bí mật")
    msg(520,1,0,"MOVE_ACCEPTED",color=GREEN,dashed=True)
    msg(650,3,1,"MOVE;RIGHT")
    msg(750,1,2,"validate + resolve")
    msg(880,2,0,"ROUND_RESULT",color=TEAL,dashed=True)
    msg(950,2,3,"ROUND_RESULT",color=TEAL,dashed=True)
    d.text((560,1040),"Không gửi lựa chọn của đối thủ trước khi RoundManager chốt kết quả.",font=font(25,True),fill="#"+RED)
    path=ASSETS/"04_sequence_lua_chon.png"; img.save(path, quality=95); return path


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + m))
        if node is None:
            node = OxmlElement("w:" + m)
            tc_mar.append(node)
        node.set(qn("w:w"), str(v)); node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=GRAY, size="8"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders"); tbl_pr.append(borders)
    for edge in ("top","left","bottom","right","insideH","insideV"):
        node = borders.find(qn("w:" + edge))
        if node is None: node = OxmlElement("w:" + edge); borders.append(node)
        node.set(qn("w:val"), "single"); node.set(qn("w:sz"), size); node.set(qn("w:color"), color)


def style_table(table, widths=None):
    table.autofit = False
    set_table_borders(table)
    for r_i,row in enumerate(table.rows):
        for c_i,cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if widths: cell.width = Cm(widths[c_i])
            set_cell_shading(cell, NAVY if r_i == 0 else (PALE_GRAY if r_i % 2 == 0 else WHITE))
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    run.font.name = "Arial"; run.font.size = Pt(9.5)
                    run.font.color.rgb = RGBColor(255,255,255) if r_i == 0 else RGBColor(17,24,39)
                    run.bold = r_i == 0


def add_table(doc, headers, rows, widths=None):
    table=doc.add_table(rows=1, cols=len(headers))
    for i,h in enumerate(headers): table.rows[0].cells[i].text=h
    for row in rows:
        cells=table.add_row().cells
        for i,value in enumerate(row): cells[i].text=str(value)
    style_table(table,widths)
    doc.add_paragraph()
    return table


def set_run_font(run, size=11, bold=False, italic=False, color=BLACK):
    run.font.name="Arial"; run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"),"Arial")
    run.font.size=Pt(size); run.bold=bold; run.italic=italic; run.font.color.rgb=RGBColor.from_string(color)


def paragraph(doc, text, bold_lead=None):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing=1.3; p.paragraph_format.space_after=Pt(6)
    if bold_lead and text.startswith(bold_lead):
        r=p.add_run(bold_lead); set_run_font(r,bold=True)
        r=p.add_run(text[len(bold_lead):]); set_run_font(r)
    else:
        r=p.add_run(text); set_run_font(r)
    return p


def bullets(doc, items):
    for item in items:
        p=doc.add_paragraph(style="List Bullet"); p.paragraph_format.space_after=Pt(3)
        set_run_font(p.add_run(item))


def heading(doc, text, level=1):
    p=doc.add_heading(text, level=level)
    p.paragraph_format.space_before=Pt(10); p.paragraph_format.space_after=Pt(5)
    for r in p.runs: set_run_font(r, size=14 if level==1 else 12, bold=True, color=BLACK)
    return p


def figure(doc, path, caption):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(6.35))
    c=doc.add_paragraph(); c.alignment=WD_ALIGN_PARAGRAPH.CENTER; c.paragraph_format.space_after=Pt(8)
    set_run_font(c.add_run(caption),size=10,bold=True)


def build_doc(images):
    doc=Document()
    sec=doc.sections[0]; sec.top_margin=Cm(2); sec.bottom_margin=Cm(2); sec.left_margin=Cm(2.3); sec.right_margin=Cm(2.3)
    normal=doc.styles["Normal"]; normal.font.name="Arial"; normal.font.size=Pt(11); normal.font.color.rgb=RGBColor.from_string(BLACK)
    title=doc.add_paragraph(); title.style=doc.styles["Title"]; title.alignment=WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(title.add_run("THÀNH VIÊN 3"),size=20,bold=True,color=BLACK)
    subtitle=doc.add_paragraph(); subtitle.alignment=WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(subtitle.add_run("Logic trọng tài và quản lý phòng chơi phía Server"),size=14,bold=True,color=NAVY)
    paragraph(doc,"Phần này trình bày phạm vi cá nhân của thành viên 3 trong hệ thống game đối kháng Client–Server. Trọng tâm là tổ chức GameRoom, điều phối sáu vòng đấu, thực thi luật thắng thua, xử lý sự cố và trả kết quả nhất quán cho hai Client.")

    heading(doc,"1. Phạm vi và nhiệm vụ",1)
    bullets(doc,[
        "Tạo GameRoom khi hai người chơi đồng ý thách đấu và quản lý vòng đời của phòng.",
        "Gán ngẫu nhiên vai Cảnh sát và Kẻ trộm; đổi vai sau mỗi vòng để mỗi người giữ từng vai đúng ba lần trong sáu vòng.",
        "Quản lý thời gian 10 giây, nhận hai lựa chọn, giữ bí mật lựa chọn và xác định người thắng vòng.",
        "Cộng điểm, xử lý timeout, Exit, mất kết nối, câu đố, Airdrop và hiệu ứng buff.",
        "Tính kết quả sau sáu vòng; nếu hòa, tiếp tục theo từng cặp hai vòng phụ và gửi kết quả cuối cho hai Client."
    ])

    heading(doc,"2. Luồng xử lý trận đấu",1)
    paragraph(doc,"MatchManager khởi tạo phòng và chuyển quyền điều phối cho GameRoom. Mỗi vòng chỉ được chốt khi đủ hai lựa chọn hợp lệ hoặc khi bộ đếm thời gian kết thúc. Sau khi công bố kết quả, vai trò được đổi và số vòng tăng lên.")
    figure(doc,images[0],"Hình 1. Luồng xử lý trận đấu do thành viên 3 phụ trách")

    heading(doc,"3. Thiết kế các lớp phía Server",1)
    paragraph(doc,"Các lớp được tách theo trách nhiệm để trạng thái phòng, luật nghiệp vụ và các tính năng phụ trợ không phụ thuộc chặt vào nhau. GameRoom giữ trạng thái trung tâm, còn các manager xử lý từng nhóm nghiệp vụ.")
    figure(doc,images[1],"Hình 2. Quan hệ giữa các lớp trong module trọng tài")
    add_table(doc,["Lớp","Trách nhiệm chính"],[
        ("GameRoom","Lưu hai người chơi, vai trò, điểm, số vòng, trạng thái và thời hạn hiện tại."),
        ("MatchManager","Tạo, tra cứu và đóng phòng; bảo đảm một người chơi chỉ thuộc một trận."),
        ("RoundManager","Bắt đầu vòng, quản lý bộ đếm 10 giây, nhận lựa chọn và chuyển vòng."),
        ("GameEngine","So sánh lựa chọn, xác định người thắng và áp dụng thay đổi điểm/buff."),
        ("PuzzleManager","Sinh câu đố, kiểm tra đáp án và ghi nhận người trả lời hợp lệ sớm nhất."),
        ("AirdropManager","Sinh vị trí Airdrop ngẫu nhiên, cấp và thu hồi hiệu ứng buff.")
    ],[3.4,12.2])

    heading(doc,"4. Trạng thái GameRoom",1)
    paragraph(doc,"Mọi thay đổi của trận đấu đi qua một trạng thái xác định. Cách tổ chức này ngăn Server nhận lựa chọn khi vòng chưa bắt đầu, chốt kết quả hai lần hoặc tiếp tục xử lý sau khi phòng đã kết thúc.")
    figure(doc,images[2],"Hình 3. Sơ đồ trạng thái của GameRoom")
    add_table(doc,["Trạng thái","Ý nghĩa","Sự kiện chuyển trạng thái"],[
        ("CREATED","Phòng đã tạo và đủ hai người chơi.","Bắt đầu chuẩn bị vòng đầu."),
        ("PREPARING_ROUND","Gán vai, làm sạch lựa chọn và đặt deadline.","Gửi ROUND_START."),
        ("WAITING_MOVES","Chờ lựa chọn hợp lệ của hai Client.","Đủ lựa chọn hoặc hết 10 giây."),
        ("RESOLVING","GameEngine tính người thắng và cập nhật điểm.","Gửi ROUND_RESULT."),
        ("ROUND_FINISHED","Kiểm tra số vòng, đổi vai hoặc xét hòa.","Vòng tiếp theo hoặc kết thúc."),
        ("FINISHED","Khóa phòng và chỉ cho phép đọc kết quả.","Giải phóng tài nguyên.")
    ],[3.2,7.0,5.4])

    heading(doc,"5. Xử lý một vòng đấu",1)
    paragraph(doc,"RoundManager chỉ chấp nhận lựa chọn đầu tiên hợp lệ của mỗi người trong thời gian cho phép. Server phản hồi xác nhận riêng cho người gửi nhưng không công bố lựa chọn của đối thủ. Khi đủ dữ liệu, GameEngine áp dụng luật và gửi cùng một kết quả đã chốt cho cả hai Client.")
    figure(doc,images[3],"Hình 4. Trình tự nhận lựa chọn và công bố kết quả")
    add_table(doc,["Điều kiện","Kết quả","Điểm"],[
        ("Police == Thief","Cảnh sát thắng vòng","Cộng điểm cho Cảnh sát"),
        ("Police != Thief","Kẻ trộm thắng vòng","Cộng điểm cho Kẻ trộm"),
        ("Một người hết 10 giây","Xử lý timeout theo luật của trận","Không nhận lựa chọn gửi muộn"),
        ("Exit hoặc mất kết nối","Kết thúc sớm và thông báo đối thủ","Khóa GameRoom")
    ],[5.0,6.5,4.1])

    heading(doc,"6. Câu đố Airdrop và buff",1)
    paragraph(doc,"PuzzleManager tạo câu đố kèm mã vòng và thời hạn. Đáp án chỉ được công nhận nếu đúng, thuộc vòng hiện tại và đến Server sớm nhất. AirdropManager chọn vị trí ngẫu nhiên từ danh sách hợp lệ, gắn buff cho người nhận và quản lý thời điểm hết hiệu lực. Mọi thay đổi buff được lưu trong GameRoom trước khi gửi thông báo đến Client.")

    heading(doc,"7. Kết thúc trận và hiệp phụ",1)
    paragraph(doc,"Sau sáu vòng chính, Server so sánh tổng điểm. Nếu có người dẫn điểm, MatchManager gửi kết quả cuối và đóng phòng. Nếu hòa, hai người chơi tiếp tục theo từng cặp hai vòng để bảo đảm mỗi người lần lượt giữ cả hai vai. Sau mỗi cặp, Server kiểm tra lại tổng điểm; quá trình dừng khi xuất hiện người thắng hoặc khi trận bị kết thúc do sự cố.")

    heading(doc,"8. Thông điệp Client Server liên quan",1)
    add_table(doc,["Thông điệp","Hướng","Nội dung xử lý"],[
        ("MATCH_START","Server -> Client","Mã phòng, đối thủ và vai trò ban đầu."),
        ("ROUND_START","Server -> Client","Số vòng, vai trò, thời hạn 10 giây và buff hiện tại."),
        ("MOVE","Client -> Server","Mã phòng, số vòng và lựa chọn của người chơi."),
        ("MOVE_ACCEPTED","Server -> Client","Xác nhận Server đã lưu lựa chọn hợp lệ."),
        ("ROUND_RESULT","Server -> Client","Hai lựa chọn, người thắng vòng, điểm và vai vòng sau."),
        ("PUZZLE / PUZZLE_ANSWER","Hai chiều","Câu đố, đáp án và thời điểm nhận tại Server."),
        ("AIRDROP","Server -> Client","Vị trí Airdrop, người nhận và loại buff."),
        ("MATCH_RESULT","Server -> Client","Người thắng trận, tổng điểm và lý do kết thúc."),
        ("OPPONENT_LEFT","Server -> Client","Đối thủ Exit hoặc mất kết nối.")
    ],[3.8,3.2,8.6])

    heading(doc,"9. Ánh xạ hình vẽ với chức năng cá nhân",1)
    add_table(doc,["Hình","Nội dung mô tả","Chức năng thành viên 3"],[
        ("Hình 1","Luồng từ tạo phòng đến đóng phòng.","Điều phối toàn bộ vòng đời trận đấu."),
        ("Hình 2","Quan hệ giữa GameRoom và các lớp quản lý.","Thiết kế module logic trọng tài phía Server."),
        ("Hình 3","Các trạng thái và điều kiện chuyển trạng thái.","Bảo vệ tính nhất quán của GameRoom."),
        ("Hình 4","Trình tự nhận và giữ kín lựa chọn.","Nhận MOVE, chốt vòng và gửi ROUND_RESULT.")
    ],[2.2,7.0,6.4])

    heading(doc,"10. Các trường hợp kiểm thử chính",1)
    add_table(doc,["Mã","Tình huống","Kết quả mong đợi"],[
        ("TC01","Hai lựa chọn giống nhau.","Cảnh sát thắng và điểm chỉ tăng một lần."),
        ("TC02","Hai lựa chọn khác nhau.","Kẻ trộm thắng và hai Client nhận cùng kết quả."),
        ("TC03","Một Client gửi lựa chọn sau 10 giây.","Server từ chối dữ liệu muộn và xử lý timeout."),
        ("TC04","Một Client gửi MOVE hai lần.","Chỉ lựa chọn hợp lệ đầu tiên được lưu."),
        ("TC05","Người chơi Exit hoặc mất kết nối.","Phòng chuyển FINISHED và đối thủ được thông báo."),
        ("TC06","Tỷ số hòa sau sáu vòng.","Server tạo đúng một cặp hai vòng phụ và tiếp tục xét điểm."),
        ("TC07","Hai đáp án câu đố đến gần nhau.","Dựa trên thời điểm Server nhận để chọn người thắng."),
        ("TC08","Airdrop xuất hiện.","Vị trí hợp lệ, buff đúng người và hết hiệu lực đúng thời điểm.")
    ],[1.5,6.0,8.1])
    doc.save(OUTPUT)


def html_table(headers, rows):
    head="".join(f"<th>{html.escape(h)}</th>" for h in headers)
    body="".join("<tr>"+"".join(f"<td>{html.escape(str(v))}</td>" for v in row)+"</tr>" for row in rows)
    return f"<table><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table>"


def image_data(path):
    return base64.b64encode(path.read_bytes()).decode("ascii")


def build_html(images):
    # The HTML mirrors the verified DOCX and is intended for one rich paste into Google Docs.
    def img(path, caption):
        return f'<p class="fig"><img src="data:image/png;base64,{image_data(path)}" style="width:680px;max-width:100%"></p><p class="cap">{html.escape(caption)}</p>'
    parts=["""<html><head><meta charset="utf-8"><style>
    body{font-family:Arial,sans-serif;color:#111827;font-size:11pt;line-height:1.35}
    h1{text-align:center;font-size:20pt;margin:0 0 4pt;color:#000} .sub{text-align:center;font-size:14pt;font-weight:bold;color:#123B5D;margin:0 0 18pt}
    h2{font-size:14pt;color:#000;margin:14pt 0 6pt} p{margin:0 0 7pt;text-align:justify} li{margin:0 0 4pt}
    table{border-collapse:collapse;width:100%;margin:8pt 0 12pt} th{background:#123B5D;color:white;font-weight:bold} th,td{border:1px solid #D9E0E5;padding:6pt;vertical-align:middle} tr:nth-child(odd) td{background:#F5F7F9}
    .fig{text-align:center;margin:10pt 0 2pt}.cap{text-align:center;font-size:10pt;font-weight:bold;margin:0 0 10pt}
    </style></head><body><h1>THÀNH VIÊN 3</h1><p class="sub">Logic trọng tài và quản lý phòng chơi phía Server</p>
    <p>Phần này trình bày phạm vi cá nhân của thành viên 3 trong hệ thống game đối kháng Client–Server. Trọng tâm là tổ chức GameRoom, điều phối sáu vòng đấu, thực thi luật thắng thua, xử lý sự cố và trả kết quả nhất quán cho hai Client.</p>
    <h2>1. Phạm vi và nhiệm vụ</h2><ul>
    <li>Tạo GameRoom khi hai người chơi đồng ý thách đấu và quản lý vòng đời của phòng.</li>
    <li>Gán ngẫu nhiên vai Cảnh sát và Kẻ trộm; đổi vai sau mỗi vòng để mỗi người giữ từng vai đúng ba lần trong sáu vòng.</li>
    <li>Quản lý thời gian 10 giây, nhận hai lựa chọn, giữ bí mật lựa chọn và xác định người thắng vòng.</li>
    <li>Cộng điểm, xử lý timeout, Exit, mất kết nối, câu đố, Airdrop và hiệu ứng buff.</li>
    <li>Tính kết quả sau sáu vòng; nếu hòa, tiếp tục theo từng cặp hai vòng phụ và gửi kết quả cuối cho hai Client.</li></ul>
    <h2>2. Luồng xử lý trận đấu</h2><p>MatchManager khởi tạo phòng và chuyển quyền điều phối cho GameRoom. Mỗi vòng chỉ được chốt khi đủ hai lựa chọn hợp lệ hoặc khi bộ đếm thời gian kết thúc. Sau khi công bố kết quả, vai trò được đổi và số vòng tăng lên.</p>""",
    img(images[0],"Hình 1. Luồng xử lý trận đấu do thành viên 3 phụ trách"),
    "<h2>3. Thiết kế các lớp phía Server</h2><p>Các lớp được tách theo trách nhiệm để trạng thái phòng, luật nghiệp vụ và các tính năng phụ trợ không phụ thuộc chặt vào nhau. GameRoom giữ trạng thái trung tâm, còn các manager xử lý từng nhóm nghiệp vụ.</p>",
    img(images[1],"Hình 2. Quan hệ giữa các lớp trong module trọng tài"),
    html_table(["Lớp","Trách nhiệm chính"],[("GameRoom","Lưu hai người chơi, vai trò, điểm, số vòng, trạng thái và thời hạn hiện tại."),("MatchManager","Tạo, tra cứu và đóng phòng; bảo đảm một người chơi chỉ thuộc một trận."),("RoundManager","Bắt đầu vòng, quản lý bộ đếm 10 giây, nhận lựa chọn và chuyển vòng."),("GameEngine","So sánh lựa chọn, xác định người thắng và áp dụng thay đổi điểm/buff."),("PuzzleManager","Sinh câu đố, kiểm tra đáp án và ghi nhận người trả lời hợp lệ sớm nhất."),("AirdropManager","Sinh vị trí Airdrop ngẫu nhiên, cấp và thu hồi hiệu ứng buff.")]),
    "<h2>4. Trạng thái GameRoom</h2><p>Mọi thay đổi của trận đấu đi qua một trạng thái xác định. Cách tổ chức này ngăn Server nhận lựa chọn khi vòng chưa bắt đầu, chốt kết quả hai lần hoặc tiếp tục xử lý sau khi phòng đã kết thúc.</p>",
    img(images[2],"Hình 3. Sơ đồ trạng thái của GameRoom"),
    html_table(["Trạng thái","Ý nghĩa","Sự kiện chuyển trạng thái"],[("CREATED","Phòng đã tạo và đủ hai người chơi.","Bắt đầu chuẩn bị vòng đầu."),("PREPARING_ROUND","Gán vai, làm sạch lựa chọn và đặt deadline.","Gửi ROUND_START."),("WAITING_MOVES","Chờ lựa chọn hợp lệ của hai Client.","Đủ lựa chọn hoặc hết 10 giây."),("RESOLVING","GameEngine tính người thắng và cập nhật điểm.","Gửi ROUND_RESULT."),("ROUND_FINISHED","Kiểm tra số vòng, đổi vai hoặc xét hòa.","Vòng tiếp theo hoặc kết thúc."),("FINISHED","Khóa phòng và chỉ cho phép đọc kết quả.","Giải phóng tài nguyên.")]),
    "<h2>5. Xử lý một vòng đấu</h2><p>RoundManager chỉ chấp nhận lựa chọn đầu tiên hợp lệ của mỗi người trong thời gian cho phép. Server phản hồi xác nhận riêng cho người gửi nhưng không công bố lựa chọn của đối thủ. Khi đủ dữ liệu, GameEngine áp dụng luật và gửi cùng một kết quả đã chốt cho cả hai Client.</p>",
    img(images[3],"Hình 4. Trình tự nhận lựa chọn và công bố kết quả"),
    html_table(["Điều kiện","Kết quả","Điểm"],[("Police == Thief","Cảnh sát thắng vòng","Cộng điểm cho Cảnh sát"),("Police != Thief","Kẻ trộm thắng vòng","Cộng điểm cho Kẻ trộm"),("Một người hết 10 giây","Xử lý timeout theo luật của trận","Không nhận lựa chọn gửi muộn"),("Exit hoặc mất kết nối","Kết thúc sớm và thông báo đối thủ","Khóa GameRoom")]),
    "<h2>6. Câu đố Airdrop và buff</h2><p>PuzzleManager tạo câu đố kèm mã vòng và thời hạn. Đáp án chỉ được công nhận nếu đúng, thuộc vòng hiện tại và đến Server sớm nhất. AirdropManager chọn vị trí ngẫu nhiên từ danh sách hợp lệ, gắn buff cho người nhận và quản lý thời điểm hết hiệu lực. Mọi thay đổi buff được lưu trong GameRoom trước khi gửi thông báo đến Client.</p>",
    "<h2>7. Kết thúc trận và hiệp phụ</h2><p>Sau sáu vòng chính, Server so sánh tổng điểm. Nếu có người dẫn điểm, MatchManager gửi kết quả cuối và đóng phòng. Nếu hòa, hai người chơi tiếp tục theo từng cặp hai vòng để bảo đảm mỗi người lần lượt giữ cả hai vai. Sau mỗi cặp, Server kiểm tra lại tổng điểm; quá trình dừng khi xuất hiện người thắng hoặc khi trận bị kết thúc do sự cố.</p>",
    "<h2>8. Thông điệp Client Server liên quan</h2>",
    html_table(["Thông điệp","Hướng","Nội dung xử lý"],[("MATCH_START","Server -> Client","Mã phòng, đối thủ và vai trò ban đầu."),("ROUND_START","Server -> Client","Số vòng, vai trò, thời hạn 10 giây và buff hiện tại."),("MOVE","Client -> Server","Mã phòng, số vòng và lựa chọn của người chơi."),("MOVE_ACCEPTED","Server -> Client","Xác nhận Server đã lưu lựa chọn hợp lệ."),("ROUND_RESULT","Server -> Client","Hai lựa chọn, người thắng vòng, điểm và vai vòng sau."),("PUZZLE / PUZZLE_ANSWER","Hai chiều","Câu đố, đáp án và thời điểm nhận tại Server."),("AIRDROP","Server -> Client","Vị trí Airdrop, người nhận và loại buff."),("MATCH_RESULT","Server -> Client","Người thắng trận, tổng điểm và lý do kết thúc."),("OPPONENT_LEFT","Server -> Client","Đối thủ Exit hoặc mất kết nối.")]),
    "<h2>9. Ánh xạ hình vẽ với chức năng cá nhân</h2>",
    html_table(["Hình","Nội dung mô tả","Chức năng thành viên 3"],[("Hình 1","Luồng từ tạo phòng đến đóng phòng.","Điều phối toàn bộ vòng đời trận đấu."),("Hình 2","Quan hệ giữa GameRoom và các lớp quản lý.","Thiết kế module logic trọng tài phía Server."),("Hình 3","Các trạng thái và điều kiện chuyển trạng thái.","Bảo vệ tính nhất quán của GameRoom."),("Hình 4","Trình tự nhận và giữ kín lựa chọn.","Nhận MOVE, chốt vòng và gửi ROUND_RESULT.")]),
    "<h2>10. Các trường hợp kiểm thử chính</h2>",
    html_table(["Mã","Tình huống","Kết quả mong đợi"],[("TC01","Hai lựa chọn giống nhau.","Cảnh sát thắng và điểm chỉ tăng một lần."),("TC02","Hai lựa chọn khác nhau.","Kẻ trộm thắng và hai Client nhận cùng kết quả."),("TC03","Một Client gửi lựa chọn sau 10 giây.","Server từ chối dữ liệu muộn và xử lý timeout."),("TC04","Một Client gửi MOVE hai lần.","Chỉ lựa chọn hợp lệ đầu tiên được lưu."),("TC05","Người chơi Exit hoặc mất kết nối.","Phòng chuyển FINISHED và đối thủ được thông báo."),("TC06","Tỷ số hòa sau sáu vòng.","Server tạo đúng một cặp hai vòng phụ và tiếp tục xét điểm."),("TC07","Hai đáp án câu đố đến gần nhau.","Dựa trên thời điểm Server nhận để chọn người thắng."),("TC08","Airdrop xuất hiện.","Vị trí hợp lệ, buff đúng người và hết hiệu lực đúng thời điểm.")]),
    "</body></html>"]
    HTML_OUTPUT.write_text("".join(parts),encoding="utf-8")


if __name__ == "__main__":
    images=[save_flow(),save_classes(),save_states(),save_sequence()]
    build_doc(images)
    build_html(images)
    print(OUTPUT)
    print(HTML_OUTPUT)
    for image in images: print(image)
